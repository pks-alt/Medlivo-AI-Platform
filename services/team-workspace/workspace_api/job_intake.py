from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import re
from typing import Any
from openpyxl import load_workbook

# Canonical Phase 1 fields. Customer spreadsheets can use any headers; mappings
# translate them into these names before validation and JobDiva write-back.
CANONICAL_FIELDS = {
    "requisition_id", "title", "facility", "city", "state", "bill_rate",
    "start_date", "end_date", "setting", "hours_per_week", "specialty",
    "schedule", "coverage_type", "providers_needed", "call_details", "notes",
}

HEADER_ALIASES = {
    "requisition id": "requisition_id",
    "req id": "requisition_id",
    "job id": "requisition_id",
    "posting title": "title",
    "job title": "title",
    "position": "title",
    "location name linked": "facility",
    "facility": "facility",
    "facility name": "facility",
    "client facility": "facility",
    "facility location city": "city",
    "city": "city",
    "facility location state province": "state",
    "state": "state",
    "state province": "state",
    "bill rate": "bill_rate",
    "start date": "start_date",
    "end date": "end_date",
    "location location setting": "setting",
    "setting": "setting",
    "hours per week": "hours_per_week",
    "weekly hours": "hours_per_week",
    "specialty": "specialty",
    "shift": "schedule",
    "schedule": "schedule",
    "coverage": "coverage_type",
    "coverage type": "coverage_type",
    "providers needed": "providers_needed",
    "fte": "providers_needed",
    "call": "call_details",
    "call details": "call_details",
    "notes": "notes",
}


def header_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", value.strip().lower()).strip()


def suggest_mapping(headers: list[str]) -> dict[str, str]:
    result: dict[str, str] = {}
    for header in headers:
        canonical = HEADER_ALIASES.get(header_key(header))
        if canonical:
            result[header] = canonical
    return result


def json_safe(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def normalize_row(source: dict[str, Any], mapping: dict[str, str]) -> dict[str, Any]:
    normalized: dict[str, Any] = {}
    for source_field, canonical_field in mapping.items():
        if canonical_field not in CANONICAL_FIELDS:
            continue
        value = source.get(source_field)
        if value is None:
            continue
        if isinstance(value, str):
            value = value.strip()
            if not value:
                continue
        normalized[canonical_field] = json_safe(value)
    return normalized


def validate_job(job: dict[str, Any], division: str) -> list[dict[str, str]]:
    errors: list[dict[str, str]] = []
    if not job.get("title"):
        errors.append({"field": "title", "code": "required", "message": "Position title is required"})
    # State is expected for all current Medlivo staffing divisions. Remote/hybrid
    # roles still need the licensing/work state represented.
    if not job.get("state"):
        errors.append({"field": "state", "code": "required", "message": "State is required"})
    if division in {"Rehabilitation", "Nursing & Allied"} and not job.get("start_date"):
        errors.append({"field": "start_date", "code": "review", "message": "Confirm assignment start date"})
    if division == "Rehabilitation" and job.get("hours_per_week") is None:
        errors.append({"field": "hours_per_week", "code": "review", "message": "Confirm weekly hours"})
    return errors


def duplicate_key(job: dict[str, Any]) -> tuple:
    if job.get("requisition_id"):
        return ("req", str(job["requisition_id"]).strip().lower())
    return (
        "composite",
        str(job.get("title", "")).strip().lower(),
        str(job.get("facility", "")).strip().lower(),
        str(job.get("city", "")).strip().lower(),
        str(job.get("state", "")).strip().lower(),
        str(job.get("start_date", "")).strip().lower(),
        str(job.get("end_date", "")).strip().lower(),
    )


def process_rows(rows: list[dict[str, Any]], mapping: dict[str, str], division: str) -> list[dict[str, Any]]:
    processed: list[dict[str, Any]] = []
    seen: set[tuple] = set()
    for index, source in enumerate(rows, start=1):
        normalized = normalize_row(source, mapping)
        errors = validate_job(normalized, division)
        key = duplicate_key(normalized)
        is_duplicate = key in seen and any(str(part) for part in key[1:])
        if not is_duplicate:
            seen.add(key)
        status = "duplicate" if is_duplicate else ("review" if errors else "ready")
        processed.append({
            "row_number": index,
            "source_row": {k: json_safe(v) for k, v in source.items()},
            "normalized_job": normalized,
            "validation_errors": errors,
            "status": status,
        })
    return processed


def parse_xlsx(content: bytes, *, max_rows: int = 500, max_columns: int = 100) -> tuple[list[str], list[dict[str, Any]]]:
    if len(content) > 5 * 1024 * 1024:
        raise ValueError("Spreadsheet exceeds the 5 MB upload limit")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header_values = next(rows)
        except StopIteration:
            raise ValueError("Spreadsheet is empty") from None
        if len(header_values) > max_columns:
            raise ValueError("Spreadsheet has too many columns")
        headers: list[str] = []
        used: set[str] = set()
        for index, value in enumerate(header_values, start=1):
            base = str(value).strip() if value is not None else f"Column {index}"
            if not base:
                base = f"Column {index}"
            candidate = base
            suffix = 2
            while candidate in used:
                candidate = f"{base} ({suffix})"
                suffix += 1
            used.add(candidate)
            headers.append(candidate)

        result: list[dict[str, Any]] = []
        for row_index, values in enumerate(rows, start=2):
            if len(result) >= max_rows:
                raise ValueError(f"Spreadsheet exceeds the {max_rows} row intake limit")
            if all(value is None or (isinstance(value, str) and not value.strip()) for value in values):
                continue
            source: dict[str, Any] = {}
            for i, header in enumerate(headers):
                value = values[i] if i < len(values) else None
                source[header] = json_safe(value)
            result.append(source)
        if not result:
            raise ValueError("Spreadsheet has no job rows")
        return headers, result
    finally:
        workbook.close()
